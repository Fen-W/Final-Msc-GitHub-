# 
#this script checks
# How much of the evidence I expected did the system actually find?
# Does the system favour long documents over short ones?
# Can BM25 keyword search tell anything apart in this corpus?
#


import re
import math
import statistics
import chromadb
import openpyxl


BASE = "/Users/woolf/Desktop/dissertation"
DATABASE = BASE + "/build_v5/vectorised_database_v5"
COLLECTION = "microplastics_corpus_v5"
WORKBOOK = BASE + "/testing/RAG_System_Testing_Workbook.xlsx"

# the words that actually appear in my twenty test questions
WORDS = ["microplastics", "atmospheric", "health", "evidence", "policy",
         "sources", "deposition", "transport", "uk", "tyre", "drinking"]


collection = chromadb.PersistentClient(path=DATABASE).get_collection(COLLECTION)
data = collection.get(include=["documents", "metadatas"])

passages = data["documents"]
total = len(passages)

# how many passages each document was cut into
size = {}
for detail in data["metadatas"]:
    name = detail["doc_id"]
    size[name] = size.get(name, 0) + 1


# pull the expected and retrieved document lists straight out of my workbook
sheet = openpyxl.load_workbook(WORKBOOK, data_only=True)["System 1"]


def document_names(cell):
    # my expected lists are typed by hand so the spacing and capitals vary
    found = re.findall(r"DOC\s?0*(\d{1,3})", str(cell or ""), re.IGNORECASE)
    names = set()
    for number in found:
        names.add("DOC%03d" % int(number))
    return names


print("")
print("1. HOW MUCH OF THE EXPECTED EVIDENCE WAS FOUND")
print("")
print("%-6s %9s %9s %6s %9s" % ("question", "expected", "returned", "hit", "reachable"))
print("-" * 45)

hits = 0
reachable = 0
returned_total = 0

for row in range(6, 26):

    expected = document_names(sheet.cell(row, 7).value)
    returned = document_names(sheet.cell(row, 11).value)

    # the out of scope questions have no expected list, so skip them
    if len(expected) == 0:
        continue

    matched = expected & returned

    # the system only ever returns six documents, so it cannot find ten.
    # the fair total is whichever is smaller.
    ceiling = min(len(expected), len(returned))

    hits = hits + len(matched)
    reachable = reachable + ceiling
    returned_total = returned_total + len(returned)

    print("%-6s %9d %9d %6d %9d" % (sheet.cell(row, 1).value,
                                    len(expected), len(returned),
                                    len(matched), ceiling))

print("-" * 45)
print("recall    %d of %d reachable = %.0f%%" % (hits, reachable, 100 * hits / reachable))
print("precision %d of %d returned  = %.0f%%" % (hits, returned_total, 100 * hits / returned_total))


print("")
print("")
print("2. DOES THE SYSTEM FAVOUR LONG DOCUMENTS")
print("")

# count how often each document was returned, and which expected ones never were
was_returned = {}
was_expected = {}

for row in range(6, 26):
    for name in document_names(sheet.cell(row, 11).value):
        was_returned[name] = was_returned.get(name, 0) + 1
    for name in document_names(sheet.cell(row, 7).value):
        was_expected[name] = was_expected.get(name, 0) + 1

everything = list(size.values())

returned_sizes = []
for name in was_returned:
    for i in range(was_returned[name]):
        returned_sizes.append(size[name])

missed_sizes = []
for name in was_expected:
    if name not in was_returned and name in size:
        missed_sizes.append(size[name])

print("median passages per document, whole corpus       %6.0f" % statistics.median(everything))
print("median passages per document, actually returned  %6.0f" % statistics.median(returned_sizes))
print("median passages per document, expected but never returned %.0f" % statistics.median(missed_sizes))
print("")
print("most returned documents:")

order = sorted(was_returned, key=lambda n: was_returned[n], reverse=True)

for name in order[:5]:
    print("   %s returned for %2d questions, %4d passages, expected for %d" %
          (name, was_returned[name], size[name], was_expected.get(name, 0)))


print("")
print("")
print("3. CAN BM25 KEYWORD SEARCH TELL ANYTHING APART")
print("")

# # BM25 favours rare words because common ones cannot separate good passages from bad. idf measures that.
words_in_passage = []
for text in passages:
    words_in_passage.append(set(re.findall(r"\b[\w-]+\b", text.lower())))

print("%-16s %10s %9s %8s" % ("word", "passages", "share", "idf"))
print("-" * 48)

for word in WORDS:

    found = 0
    for passage in words_in_passage:
        if word in passage:
            found = found + 1

    idf = math.log(1 + (total - found + 0.5) / (found + 0.5))

    note = ""
    if idf < 2:
        note = "  too common to be useful"

    print("%-16s %10d %8.1f%% %8.2f%s" % (word, found, 100 * found / total, idf, note))

print("")
print("total passages searched:", total)
print("")
